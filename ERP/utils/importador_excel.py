# utils/importador_excel.py
# Importa datos desde Excel a cualquier tabla de la BBDD del ERP

import sys
import openpyxl
import mysql.connector

# 🔧 Configuración base de datos (ajusta según tu entorno)
DB_CONFIG = {
    'unix_socket': '/Applications/MAMP/tmp/mysql/mysql.sock',
    'user': 'horti',
    'password': 'horti',
    'database': 'erp'
}

def conectar_bbdd():
    """Establece la conexión a la base de datos"""
    try:
        conn = mysql.connector.connect(**DB_CONFIG)
        return conn, conn.cursor()
    except mysql.connector.Error as err:
        print(f"❌ Error al conectar con la BBDD: {err}")
        sys.exit(1)

def leer_excel(archivo):
    """Lee el archivo Excel y devuelve cabeceras y filas"""
    try:
        wb = openpyxl.load_workbook(archivo)
        hoja = wb.active
        cabeceras = [str(cell.value).strip() if cell.value else '' for cell in hoja[1]]
        filas = []
        for fila in hoja.iter_rows(min_row=2, values_only=True):
            filas.append(tuple(cell if cell != '' else None for cell in fila))
        return cabeceras, filas
    except Exception as e:
        print(f"❌ Error leyendo el Excel: {e}")
        sys.exit(1)

def insertar_datos(tabla, cabeceras, filas):
    """Inserta los datos fila por fila en la tabla especificada"""
    conn, cursor = conectar_bbdd()

    placeholders = ', '.join(['%s'] * len(cabeceras))
    columnas = ', '.join([f"`{col}`" for col in cabeceras])
    sql = f"INSERT INTO `{tabla}` ({columnas}) VALUES ({placeholders})"

    insertados, errores = 0, 0

    for i, fila in enumerate(filas):
        try:
            cursor.execute(sql, fila)
            insertados += 1
        except mysql.connector.Error as err:
            errores += 1
            print(f"⚠️ Fila {i + 2} no insertada: {err.msg}")

    conn.commit()
    cursor.close()
    conn.close()

    print(f"\n✅ Insertados: {insertados}")
    if errores:
        print(f"❌ Fallos: {errores}")

def main():
    if len(sys.argv) != 3:
        print("Uso: python3 importador_excel.py <tabla_destino> <archivo_excel>")
        sys.exit(1)

    tabla = sys.argv[1]
    archivo = sys.argv[2]

    print(f"📥 Archivo: {archivo} → Tabla: {tabla}")
    cabeceras, filas = leer_excel(archivo)

    if not cabeceras or not filas:
        print("❌ Archivo vacío o mal formado.")
        sys.exit(1)

    print(f"📊 Columnas: {cabeceras}")
    print(f"📄 Registros detectados: {len(filas)}")

    confirmar = input("¿Deseas continuar con la importación? (s/n): ").lower()
    if confirmar != 's':
        print("⛔ Importación cancelada.")
        sys.exit(0)

    insertar_datos(tabla, cabeceras, filas)

if __name__ == '__main__':
    main()
